from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Security, status, BackgroundTasks, UploadFile, File, Form
from fastapi.security import APIKeyHeader
from fastapi.responses import FileResponse
import docx
import io
import os
import base64
from pydantic import BaseModel

from scrapers.scraper import scrape_all
from utils.logger import log_error
from utils.gemini_helper import generate_role_summary

from scrapers.scraper_workflow import run_scraper_workflow

load_dotenv()

app = FastAPI(title="UTSA Digital Tools Scraper API")

API_KEY_NAME = "X-API-KEY"
api_key_header = APIKeyHeader(name=API_KEY_NAME, auto_error=False)


def verify_api_key(api_key: str | None = Security(api_key_header)) -> str:
    expected_key = os.getenv("SCRAPER_API_KEY")
    if not expected_key or api_key != expected_key:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid or missing API key",
        )
    return api_key

@app.get("/")
def read_root():
    return {"status": "online", "service": "UTSA Scraper API"}

class FormatRequest(BaseModel):
    file_base64: str
    user_group: str

@app.post("/format-release-notes")
async def format_release_notes(request: FormatRequest):
    try:
        file_bytes = base64.b64decode(request.file_base64)
        doc = docx.Document(io.BytesIO(file_bytes))
        
        # Extraer párrafos
        text_runs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Extraer tablas
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_runs.append(" | ".join(row_text))
                        
        raw_text = "\n\n".join(text_runs)

        if not raw_text.strip():
            return {
                "status": "warning",
                "user_group": request.user_group,
                "formatted_content": "No readable text found in the uploaded document."
            }
        
        # Generar el resumen en Markdown mediante Gemini
        formatted_text = generate_role_summary(raw_text, request.user_group)
        
        return {
            "status": "success",
            "user_group": request.user_group,
            "formatted_content": formatted_text  # Devuelve el texto Markdown directo
        }
    except Exception as e:
        log_error(f"Error processing release notes: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "healthy"}

@app.post("/run-scraper")
def trigger_scraper():
    doc_path = run_scraper_workflow()
    
    if not doc_path or not os.path.exists(doc_path):
        raise HTTPException(status_code=500, detail="Document generation failed or returned invalid path.")
        
    return FileResponse(
        path=doc_path, 
        filename=os.path.basename(doc_path),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

@app.post("/scrape")
def trigger_scrape(_: str = Security(verify_api_key)) -> dict:
    try:
        raw_results = scrape_all()

        current_results = []
        for item in raw_results:
            url = item.get("url")
            if not url:
                continue

            content = item.get("content", "")
            if any(
                error_text in content
                for error_text in (
                    "ERR_HTTP2_PROTOCOL_ERROR",
                    "This site can't be reached",
                )
            ):
                continue

            current_results.append(item)

        return {
            "status": "success",
            "count": len(current_results),
            "data": current_results,
        }
    except Exception as exc:
        log_error(f"API scrape failed: {exc}")
        raise HTTPException(status_code=500, detail="Scrape failed") from exc