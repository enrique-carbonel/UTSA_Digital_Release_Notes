import os
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from utils.logger import log_info

def create_release_notes_doc(notes, summary):
    date_str = datetime.now().strftime("%B %d, %Y")
    filename = f"UTSA_Digital_Tools_Release_Notes_{datetime.now().strftime('%Y%m%d')}.docx"
    output_path = Path(__file__).resolve().parent.parent / "RAW" / filename
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    doc = Document()
    
    # Title
    title = doc.add_heading(f"UTSA Digital Tools Release Notes - {date_str}", 0)
    title.runs[0].font.color.rgb = RGBColor(12, 35, 64)
    
    # AI Summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(summary)
    
    # Individual Tools
    doc.add_heading("Detailed Tool Updates", level=1)
    
    for note in notes:
        h2 = doc.add_heading(note['tool'], level=2)
        doc.add_paragraph(f"Source: {note['url']}", style='Intense Quote')
            
        # NEW: Check if this tool has a structured table
        if note.get('is_table'):
            table = doc.add_table(rows=0, cols=len(note['data'][0]))
            table.style = 'Table Grid'
            for row_data in note['data']:
                row_cells = table.add_row().cells
                for i, text in enumerate(row_data):
                    row_cells[i].text = text
        else:
            doc.add_heading(note['title'], level=3)
            doc.add_paragraph(note['content'])
                
        doc.add_paragraph("-" * 40)
        
    doc.save(output_path)
    log_info(f"Document saved as {output_path}")
    return str(output_path)